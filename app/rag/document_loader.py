"""文档加载器：读取 ai_data 目录下的 PDF / DOCX 教学文档。

输出按页组织的纯文本，保留页码以便溯源。
"""
from __future__ import annotations

from dataclasses import dataclass, field
from pathlib import Path

SUPPORTED_EXTS = {".pdf", ".docx"}


@dataclass
class RawDocument:
    path: Path
    title: str          # 文件名（不含扩展名）
    file_type: str      # pdf / docx
    pages: list[tuple[int, str]] = field(default_factory=list)  # [(page_no, text)]


def load_file(path: str | Path) -> RawDocument:
    """加载单个 pdf/docx 文件为按页文本。"""
    p = Path(path)
    if p.suffix.lower() not in SUPPORTED_EXTS:
        raise ValueError(f"不支持的文件类型: {p.suffix}")

    title = p.stem
    file_type = p.suffix.lower().lstrip(".")

    if file_type == "pdf":
        pages = _load_pdf(p)
    else:
        pages = _load_docx(p)
    return RawDocument(path=p, title=title, file_type=file_type, pages=pages)


def load_dir(directory: str | Path) -> list[RawDocument]:
    """加载目录下所有受支持的文档。"""
    folder = Path(directory)
    docs = []
    for p in sorted(folder.iterdir()):
        if p.is_file() and p.suffix.lower() in SUPPORTED_EXTS:
            docs.append(load_file(p))
    return docs


def _load_pdf(path: Path) -> list[tuple[int, str]]:
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    pages = []
    for i, page in enumerate(reader.pages, start=1):
        text = page.extract_text() or ""
        if text.strip():
            pages.append((i, text.strip()))
    return pages


def _load_docx(path: Path) -> list[tuple[int, str]]:
    from docx import Document as DocxDocument

    doc = DocxDocument(str(path))
    parts = [p.text for p in doc.paragraphs if p.text.strip()]
    tables = []
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells]
            tables.append(" | ".join(cells))
    text = "\n".join([*parts, *tables])
    return [(0, text)] if text.strip() else []
